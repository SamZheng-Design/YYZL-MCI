#!/usr/bin/env python3
"""
生成中流通性能优化实施 Prompts 文档 (Word格式)
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import datetime

doc = Document()

# ── 全局样式设置 ──
style = doc.styles['Normal']
font = style.font
font.name = 'Microsoft YaHei'
font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.3

# 标题样式
for level in range(1, 4):
    hs = doc.styles[f'Heading {level}']
    hs.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)
    hs.font.bold = True

# ── 封面 ──
doc.add_paragraph()
doc.add_paragraph()
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('中流通 ZhongLiu Connect')
run.font.size = Pt(28)
run.font.bold = True
run.font.color.rgb = RGBColor(0xb9, 0x1c, 0x1c)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('全站性能优化实施 Prompts')
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(0x4a, 0x4a, 0x4a)

doc.add_paragraph()
info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = info.add_run(f'生成日期: {datetime.date.today()}\n适用于新对话窗口，按顺序逐段发送')
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.add_page_break()

# ── 使用说明 ──
doc.add_heading('使用说明', level=1)
doc.add_paragraph(
    '本文档包含 8 个 Prompt，请在新对话中按顺序发送。\n'
    '每个 Prompt 是一个独立的实施步骤，等上一步完成后再发下一个。\n'
    '如果某一步遇到问题，可以在当前对话中继续沟通解决，不需要重新开始。'
)

usage_items = [
    'Prompt 0 — 项目上下文初始化（必须最先发送）',
    'Prompt 1 — Phase 1A: 数据库层优化（N+1修复 + 新索引）',
    'Prompt 2 — Phase 1B: API 化改造（首页 + 项目列表）',
    'Prompt 3 — Phase 1C: API 化改造（还款页 + 教师页 + 管理后台）',
    'Prompt 4 — Phase 2A: 客户端路由（SPA化 + 骨架屏过渡）',
    'Prompt 5 — Phase 2B: 分页 + 按需加载',
    'Prompt 6 — Phase 3: 部署到 zhongliutong.net + 性能验证',
    'Prompt 7 — Phase 3+: 监控 + 安全加固 + 上线检查清单',
]
for item in usage_items:
    p = doc.add_paragraph(item, style='List Bullet')

doc.add_paragraph()
p = doc.add_paragraph()
run = p.add_run('重要提示：')
run.font.bold = True
run.font.color.rgb = RGBColor(0xb9, 0x1c, 0x1c)
p.add_run(' Prompt 0 是项目完整上下文，篇幅较长但必须完整发送，它让新对话了解整个项目的架构、文件结构、当前状态和性能数据。')

doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# PROMPT 0 — 项目上下文
# ══════════════════════════════════════════════════════════════
doc.add_heading('Prompt 0 — 项目上下文初始化', level=1)
doc.add_paragraph('复制以下灰色框内全部内容，粘贴到新对话的第一条消息中：')

prompt0 = '''你好，我是 Sam，我正在开发一个叫「中流通 ZhongLiu Connect」的金融教育平台。项目已经部署到 zhongliutong.net，现在需要做全站性能优化。1000多个学生和30个老师要用，当前页面跳转很慢。

请先阅读以下项目完整上下文，不要做任何修改，等我后续发具体的优化任务。

## 技术架构
- **框架**: Hono (TypeScript) + Cloudflare Pages
- **数据库**: Cloudflare D1 (SQLite)
- **前端**: 纯 HTML/CSS/JS（无前端框架），Tailwind CSS v4，CDN 加载 FontAwesome
- **部署**: Cloudflare Pages，wrangler CLI
- **项目目录**: /home/user/webapp
- **Cloudflare 项目名**: zhongliu-connect
- **线上地址**: https://zhongliutong.net
- **D1 数据库名**: zhongliu-production (ID: 887cd767-4078-4332-83cc-e1b4f0531673)

## 核心问题（当前性能数据）
页面 TTFB 和 HTML 大小：
| 页面 | TTFB | HTML大小 | DB查询数 |
|------|------|----------|----------|
| /login | 134ms | 134KB | 0 |
| / (首页) | 643ms | 161KB | 6 |
| /projects | 880ms | 143KB | 6 |
| /repayments | 696ms | 812KB | 6+ |
| /teacher | 833ms | 136KB | 4+ |
| /admin | 577ms | 304KB | 5+ |
| /create | 139ms | 178KB | 0 |
| /profile | 126ms | 113KB | 1 |
| /notifications | 118ms | 98KB | 0 |
| /investments | 125ms | 97KB | 0 |

**根因**: 全量 SSR 架构 — 每个页面在 Worker 端执行 4-6 次 D1 查询，将全部数据通过 JSON.stringify 嵌入 HTML 返回。没有 SPA 路由，每次导航都是完整页面重载。

## 项目文件结构
```
webapp/
├── src/
│   ├── index.tsx           (554行) 主入口，路由注册 + API endpoints
│   ├── renderer.tsx        (378行) HTML 骨架渲染器
│   ├── db-bridge.ts        (685行) 数据库查询层（所有 load* 函数）
│   ├── types.ts            (426行) TypeScript 类型定义
│   ├── components.tsx      (933行) 共享 UI 组件
│   ├── data.ts             (1934行) 演示数据/常量
│   ├── middleware.ts       (328行) 认证中间件
│   ├── ai-assistant.tsx    AI 助手组件
│   ├── admin-api.ts        管理端 API
│   └── routes/
│       ├── admin.tsx       (1594行) 管理后台
│       ├── create.tsx      (1534行) 创建项目
│       ├── project-detail.tsx (1206行) 项目详情
│       ├── login.tsx       (793行) 登录页
│       ├── terms-connect.tsx (716行) 条款连接
│       ├── contract-sign.tsx (503行) 合同签署
│       ├── teacher.tsx     (477行) 教师端
│       ├── home.tsx        (402行) 首页
│       ├── repayments.tsx  (356行) 还款管理
│       ├── revenue-report.tsx (283行) 营收报告
│       ├── projects.tsx    (266行) 项目列表
│       ├── investments.tsx (257行) 投资记录
│       ├── profile.tsx     (227行) 个人资料
│       ├── notifications.tsx (147行) 通知
│       └── share.tsx       (42行) 分享页
├── public/static/
│   ├── app.css             (154KB) 全站样式
│   └── tailwind.css        (14KB) Tailwind 编译
├── migrations/
│   ├── 0001_initial_schema.sql
│   ├── 0002_rate_limits.sql
│   ├── 0003_contract_terms_upgrade.sql
│   └── 0004_expect_multiple.sql
├── wrangler.jsonc
├── ecosystem.config.cjs
├── package.json
├── vite.config.ts
└── tsconfig.json
```

## 关键架构细节

### 路由注册方式
每个路由文件导出一个 `registerXxxRoute(app)` 函数，在 `src/index.tsx` 中调用注册。
路由内使用 `app.get('/path', async (c) => { ... })` 返回完整 HTML。

### 数据查询层 (db-bridge.ts)
核心查询函数：
- `loadMembers(db)` — SELECT * FROM users WHERE role='member'
- `loadTeachers(db)` — SELECT * FROM users WHERE role='teacher'  
- `loadProjects(db)` — SELECT * FROM projects + **逐个查 project_investors（N+1问题）**
- `loadContracts(db)` — SELECT * FROM contracts
- `loadRepaymentRecords(db)` — SELECT * FROM repayment_details
- `loadRepayments(db)` — SELECT * FROM settlement_records
- `loadRevenueReports(db)` — SELECT * FROM settlement_records (grouped)
- `loadReferrals(db)` — SELECT * FROM referrals
- `loadNotifications(db)` — SELECT * FROM notifications
- `loadShareLogs(db)` — SELECT * FROM share_logs

### 已有 API 端点
已经存在一套 /api/data/* 端点可以返回 JSON：
- GET /api/data/members
- GET /api/data/teachers
- GET /api/data/projects
- GET /api/data/projects/:id
- GET /api/data/contracts
- GET /api/data/contracts/project/:projectId
- GET /api/data/revenue-reports
- GET /api/data/repayment-records
- GET /api/data/repayments
- GET /api/data/referrals
- GET /api/data/notifications
- GET /api/data/share-logs
还有: POST /api/login, POST /api/change-password, POST /api/self-register 等

### Worker 打包
- 当前 _worker.js: 743KB (gzip 175KB)
- 已提取 125KB 内联 CSS 到 /static/app.css（已完成）
- app.css 通过 CDN 缓存（max-age=14400）

### 数据库索引
已有完善的索引（在 0001_initial_schema.sql 中定义了 30+ 个索引）。

### PM2 本地开发配置
```javascript
// ecosystem.config.cjs
module.exports = {
  apps: [{
    name: 'zhongliu-connect',
    script: 'npx',
    args: 'wrangler pages dev dist --d1=zhongliu-production --local --ip 0.0.0.0 --port 3000',
    cwd: '/home/user/webapp',
    env: { NODE_ENV: 'development', PORT: 3000 },
    watch: false, instances: 1, exec_mode: 'fork'
  }]
}
```

## 已确认的优化方案（三阶段）

### Phase 1: 立竿见影（TTFB降80%）
1A. 数据库层：修复 loadProjects N+1 查询，用 JOIN 替代循环
1B. API化改造：首页、项目列表 — 页面只返回 HTML 骨架 + 骨架屏，数据通过 /api/ 异步加载
1C. API化改造：还款页、教师页、管理后台

### Phase 2: 用户体验质变
2A. 客户端路由 SPA 化 — 拦截 <a> 点击，fetch 加载内容，避免全页重载
2B. 分页 + 按需加载 — 列表默认 20 条，滚动加载更多

### Phase 3: 部署验证 + 监控
3A. 部署到 zhongliutong.net，性能对比验证
3B. 前端性能监控 + 错误追踪 + 上线检查清单

请确认你已理解以上项目上下文，然后我会发送具体的优化任务。'''

p = doc.add_paragraph()
p.style = doc.styles['Normal']
# Add as a styled block
run = p.add_run(prompt0)
run.font.size = Pt(10)
run.font.name = 'Consolas'
# Add border/shading via paragraph formatting
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
shading = OxmlElement('w:shd')
shading.set(qn('w:fill'), 'F5F5F5')
shading.set(qn('w:val'), 'clear')
p.paragraph_format.element.get_or_add_pPr().append(shading)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# PROMPT 1 — Phase 1A: 数据库层优化
# ══════════════════════════════════════════════════════════════
doc.add_heading('Prompt 1 — Phase 1A: 数据库层优化', level=1)
doc.add_paragraph('等 AI 确认理解上下文后，发送此 Prompt：')

prompt1 = '''开始 Phase 1A：数据库层优化。

## 任务 1: 修复 loadProjects 的 N+1 查询

当前 `src/db-bridge.ts` 中的 `loadProjects` 函数：
1. 先 `SELECT * FROM projects ORDER BY created_at DESC`
2. 然后对每个 project 单独执行 `SELECT investor_id FROM project_investors WHERE project_id = ?`
这在有 1000 个项目时会产生 1001 次 SQL 查询。

请修改为一次 JOIN 查询或两次查询（先查 projects，再一次性查所有 project_investors，然后在 JS 中合并）。

## 任务 2: 优化其他查询函数

检查 db-bridge.ts 中所有 `load*` 函数，看是否有类似的 N+1 问题或者可以优化的地方：
- 是否有 SELECT * 可以改为只查需要的字段？
- 是否有可以合并的查询？

## 任务 3: 添加复合索引（如果需要）

当前已有单列索引，检查是否需要添加复合索引来优化常用查询模式。如果需要，创建 `migrations/0005_performance_indexes.sql`。

## 要求
- 修改完后本地构建 `npm run build` 并启动服务验证
- 用 curl 测试相关 API 端点确认数据正确
- git commit 保存进度
- 告诉我优化前后的查询次数对比'''

p = doc.add_paragraph()
run = p.add_run(prompt1)
run.font.size = Pt(10)
run.font.name = 'Consolas'
shading = OxmlElement('w:shd')
shading.set(qn('w:fill'), 'F5F5F5')
shading.set(qn('w:val'), 'clear')
p.paragraph_format.element.get_or_add_pPr().append(shading)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# PROMPT 2 — Phase 1B: API 化改造（首页 + 项目列表）
# ══════════════════════════════════════════════════════════════
doc.add_heading('Prompt 2 — Phase 1B: API化改造（首页 + 项目列表）', level=1)
doc.add_paragraph('等 Phase 1A 完成并确认后，发送此 Prompt：')

prompt2 = '''开始 Phase 1B：首页和项目列表页的 API 化改造。

## 核心思路
当前问题：每个页面在 Worker 端做 4-6 次 D1 查询，把全部数据 JSON.stringify 嵌入 HTML。
目标：页面只返回 HTML 骨架 + 骨架屏（loading 状态），数据通过 /api/ 端点异步加载。

## 任务 1: 改造首页 (src/routes/home.tsx)

当前首页执行：loadMembers, loadProjects, loadRepayments, loadContracts, loadRepaymentRecords, loadTeachers（6次查询），TTFB 643ms。

改造为：
1. `registerHomeRoute` 不再执行任何 DB 查询
2. 返回的 HTML 中包含骨架屏（加载动画占位符），样式要和当前页面一致
3. 页面加载后，前端 JS 调用已有的 `/api/data/*` 端点获取数据
4. 数据到达后，用 JS 渲染替换骨架屏
5. 重要：保持现有的所有前端交互功能不变（标签切换、角色判断、项目卡片点击等）

## 任务 2: 改造项目列表页 (src/routes/projects.tsx)

和首页同样的改造思路：
1. 去掉 SSR 数据查询
2. 添加骨架屏
3. 前端异步加载

## 任务 3: 优化 /api/data/* 端点

检查 `src/index.tsx` 中现有的 /api/data/* 端点：
- 确保它们返回的数据格式满足前端需要
- 如果需要新的聚合端点（比如首页需要的统计数据），就创建
- 考虑添加简单的缓存头 (Cache-Control)

## 骨架屏设计要求
- 使用 Tailwind CSS 的 animate-pulse 动画
- 骨架颜色和当前深色/浅色主题匹配
- 骨架形状和实际内容布局一致（卡片、列表、统计数字等）
- 加载完成后平滑过渡到真实内容

## 要求
- 改造后首页和项目列表页的 TTFB 应该降到 ~130ms（和 /login 一样快）
- 所有现有功能必须正常工作
- 本地构建并测试，确认页面正常加载数据
- git commit 保存进度
- 给我看改造前后的 TTFB 对比'''

p = doc.add_paragraph()
run = p.add_run(prompt2)
run.font.size = Pt(10)
run.font.name = 'Consolas'
shading = OxmlElement('w:shd')
shading.set(qn('w:fill'), 'F5F5F5')
shading.set(qn('w:val'), 'clear')
p.paragraph_format.element.get_or_add_pPr().append(shading)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# PROMPT 3 — Phase 1C: API 化改造（还款页 + 教师页 + 管理后台）
# ══════════════════════════════════════════════════════════════
doc.add_heading('Prompt 3 — Phase 1C: API化改造（还款页 + 教师页 + 管理后台）', level=1)
doc.add_paragraph('等 Phase 1B 完成并确认后，发送此 Prompt：')

prompt3 = '''开始 Phase 1C：将剩余的高 TTFB 页面进行 API 化改造。

## 需要改造的页面（按 TTFB 严重程度排序）

### 1. /repayments (TTFB: 696ms, HTML: 812KB) — 最严重
当前加载：contracts, members, projects, repaymentRecords, revenueReports, shareLogs
812KB 的 HTML 说明把所有还款记录都内联了，这是最大的性能杀手。
- 改为骨架屏 + API 加载
- 重要：还款数据应该按月/按项目分页，不要一次性全部加载

### 2. /teacher (TTFB: 833ms, HTML: 136KB)
当前加载：members, projects, referrals, teachers
- 改为骨架屏 + API 加载

### 3. /admin (TTFB: 577ms, HTML: 304KB)
当前加载：contracts, members, projects, repaymentRecords, teachers
17次 JSON.stringify，最多的嵌入数据。
- 改为骨架屏 + API 加载

### 4. 其他需要检查的页面
- /revenue-report (TTFB 应该也较高)
- /investments
- /contract-sign
- /project-detail

## 改造方式
和 Phase 1B 一样：
1. 路由不做 DB 查询，只返回 HTML 骨架 + 骨架屏
2. 前端 JS 调用 /api/data/* 加载数据
3. 骨架屏风格和主题一致

## 要求
- 所有页面 TTFB 目标: < 200ms
- 所有现有功能保持正常
- 对于数据量大的页面（repayments），前端加载时显示进度
- 本地测试全部页面
- git commit 保存进度
- 给我一个所有页面改造前后 TTFB 的完整对比表'''

p = doc.add_paragraph()
run = p.add_run(prompt3)
run.font.size = Pt(10)
run.font.name = 'Consolas'
shading = OxmlElement('w:shd')
shading.set(qn('w:fill'), 'F5F5F5')
shading.set(qn('w:val'), 'clear')
p.paragraph_format.element.get_or_add_pPr().append(shading)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# PROMPT 4 — Phase 2A: 客户端路由
# ══════════════════════════════════════════════════════════════
doc.add_heading('Prompt 4 — Phase 2A: 客户端路由（SPA化）', level=1)
doc.add_paragraph('等 Phase 1 全部完成后，发送此 Prompt：')

prompt4 = '''开始 Phase 2A：客户端路由 SPA 化。

## 核心目标
当前每次点击导航链接，浏览器都做一次完整的页面重载（白屏 → 等待 TTFB → 重新渲染）。
改为 SPA 式导航：保留侧边栏/顶栏不变，只替换内容区域。

## 实现方案

### 1. 创建全局路由器 (public/static/router.js 或内嵌)
- 拦截所有内部 `<a>` 标签的 click 事件
- 排除外部链接、target="_blank"、/api/* 路径
- 使用 `fetch` 请求目标页面的 HTML
- 从返回的 HTML 中提取内容区域（#zlc-page-wrap 内的内容）
- 替换当前页面的内容区域
- 更新 `history.pushState` 和 document.title
- 处理浏览器前进/后退 (popstate)

### 2. 页面过渡动画
- 切换时显示顶部进度条或内容区 fade 效果
- 骨架屏加载状态

### 3. 注意事项
- 侧边栏高亮需要根据新 URL 更新
- 页面内的 `<script>` 标签需要重新执行
- 登录状态检查需要保持
- /login 页面不参与 SPA 路由（保持完整页面）
- 需要处理 `/project-detail/:id` 这类动态路由

### 4. 桌面端侧边栏和顶栏
当前桌面端布局有固定侧边栏和顶栏（在 renderer.tsx 中通过 JS 动态创建）。
SPA 路由切换时这些不应该被替换。

## 要求
- 页面切换体验应该是"即时"的（< 300ms 视觉反馈）
- 所有页面功能正常（包括表单提交、弹窗等）
- 浏览器前进/后退正常工作
- 本地测试所有主要导航路径
- git commit 保存进度'''

p = doc.add_paragraph()
run = p.add_run(prompt4)
run.font.size = Pt(10)
run.font.name = 'Consolas'
shading = OxmlElement('w:shd')
shading.set(qn('w:fill'), 'F5F5F5')
shading.set(qn('w:val'), 'clear')
p.paragraph_format.element.get_or_add_pPr().append(shading)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# PROMPT 5 — Phase 2B: 分页
# ══════════════════════════════════════════════════════════════
doc.add_heading('Prompt 5 — Phase 2B: 分页 + 按需加载', level=1)
doc.add_paragraph('等 Phase 2A 完成后，发送此 Prompt：')

prompt5 = '''开始 Phase 2B：数据分页和按需加载。

## 核心目标
当前所有列表页一次性加载全部数据。1000个学生 + 数百个项目/合同/还款记录 = 大量数据传输。
改为分页加载，每次只传 20 条。

## 任务 1: 后端分页 API

为以下端点添加分页支持：
- `GET /api/data/projects?page=1&limit=20&status=open&search=xxx`
- `GET /api/data/contracts?page=1&limit=20&project_id=xxx`
- `GET /api/data/repayment-records?page=1&limit=20&month=2026-03`
- `GET /api/data/members?page=1&limit=20&class_id=xxx&search=xxx`
- `GET /api/data/notifications?page=1&limit=20`

响应格式（已在 types.ts 中定义了 PaginatedResponse）：
```json
{
  "ok": true,
  "data": [...],
  "total": 150,
  "page": 1,
  "limit": 20,
  "total_pages": 8
}
```

在 D1 查询中使用 `LIMIT ? OFFSET ?` 和 `SELECT COUNT(*)` 来实现。

## 任务 2: 前端分页 UI

为列表页添加：
- 底部分页控件（上一页/下一页/页码）
- 或无限滚动（滚到底部自动加载下一页）
- 搜索/筛选时重置到第 1 页
- 加载中显示列表项的骨架屏

## 任务 3: 特别处理 /repayments 页面

这个页面之前 812KB，需要：
- 按月份筛选（默认当月）
- 按项目筛选
- 分页显示还款明细
- 目标：每次 API 返回 < 10KB

## 要求
- 确保分页查询正确使用索引（EXPLAIN QUERY PLAN）
- 本地测试分页功能
- git commit 保存进度'''

p = doc.add_paragraph()
run = p.add_run(prompt5)
run.font.size = Pt(10)
run.font.name = 'Consolas'
shading = OxmlElement('w:shd')
shading.set(qn('w:fill'), 'F5F5F5')
shading.set(qn('w:val'), 'clear')
p.paragraph_format.element.get_or_add_pPr().append(shading)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# PROMPT 6 — Phase 3: 部署 + 性能验证
# ══════════════════════════════════════════════════════════════
doc.add_heading('Prompt 6 — Phase 3: 部署 + 性能验证', level=1)
doc.add_paragraph('等所有优化完成后，发送此 Prompt：')

prompt6 = '''开始 Phase 3：部署到 zhongliutong.net 并做性能验证。

## 任务 1: 生产部署

1. 确认所有代码已 git commit
2. 运行 `npm run build`
3. 部署到 Cloudflare Pages production:
   ```
   npx wrangler pages deploy dist --project-name zhongliu-connect --branch main
   ```
4. 如果有数据库迁移（0005 等），先执行：
   ```
   npx wrangler d1 migrations apply zhongliu-production
   ```

## 任务 2: 性能验证

部署后，用 curl 测试所有页面的 TTFB：
```bash
for path in /login / /projects /repayments /teacher /admin /create /profile /notifications /investments; do
  echo "=== $path ==="
  curl -s -o /dev/null -w "TTFB: %{time_starttransfer}s | Total: %{time_total}s | Size: %{size_download}" https://zhongliutong.net$path
  echo ""
done
```

给我一个优化前后对比表：
| 页面 | 优化前TTFB | 优化后TTFB | 优化前HTML | 优化后HTML | 改善 |
|------|-----------|-----------|-----------|-----------|------|

## 任务 3: 浏览器端验证

用 Playwright 测试：
1. 打开首页，测量 DOMContentLoaded 和 Load 时间
2. 点击导航到 /projects，测量切换时间（SPA 路由应该 < 300ms）
3. 截图确认页面渲染正确

## 任务 4: CDN 缓存检查

确认以下静态资源已被 CDN 缓存：
- /static/app.css
- /static/tailwind.css
- /static/router.js（如果有的话）

检查 Cache-Control 头和 cf-cache-status。

## 要求
- 所有页面 TTFB < 200ms（目标）
- 所有功能正常
- 给我完整的测试报告和截图'''

p = doc.add_paragraph()
run = p.add_run(prompt6)
run.font.size = Pt(10)
run.font.name = 'Consolas'
shading = OxmlElement('w:shd')
shading.set(qn('w:fill'), 'F5F5F5')
shading.set(qn('w:val'), 'clear')
p.paragraph_format.element.get_or_add_pPr().append(shading)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# PROMPT 7 — Phase 3+: 监控 + 安全 + 上线清单
# ══════════════════════════════════════════════════════════════
doc.add_heading('Prompt 7 — Phase 3+: 监控 + 安全加固 + 上线检查清单', level=1)
doc.add_paragraph('部署验证通过后，发送此 Prompt：')

prompt7 = '''最后一步：上线前的安全加固和监控。

## 任务 1: 前端性能监控

在 renderer.tsx 中添加简单的 Web Vitals 上报：
- 采集 LCP (Largest Contentful Paint)
- 采集 FCP (First Contentful Paint)
- 采集 CLS (Cumulative Layout Shift)
- 通过 /api/metrics POST 上报到 Worker，Worker 写入 D1 或 KV

## 任务 2: API 错误处理加固

1. 所有 /api/* 端点添加 try-catch，返回统一的错误格式：
   ```json
   {"ok": false, "error": "描述信息", "code": "ERROR_CODE"}
   ```
2. 前端 API 调用添加超时（5秒）和重试（1次）
3. 网络错误时显示友好的提示（"网络连接异常，请重试"）

## 任务 3: 安全检查清单

请检查并修复：
- [ ] 所有 API 端点是否有权限验证（不能未登录就访问 /api/data/*）
- [ ] 管理员端点是否检查角色
- [ ] SQL 查询是否都用了参数化（防注入）
- [ ] 密码是否 hash 存储
- [ ] Session token 是否安全生成
- [ ] CORS 设置是否合理

## 任务 4: 上线检查清单

给我一份完整的上线前检查清单，包括：
- 数据库备份方案
- 回滚方案
- 监控告警
- 容量评估（1000学生 + 30老师的 D1 QPS 预估）
- 需要告知学生/老师的注意事项

## 要求
- 实施所有安全加固
- 生成检查清单文档
- 最终部署到 zhongliutong.net
- git commit + push 到 GitHub'''

p = doc.add_paragraph()
run = p.add_run(prompt7)
run.font.size = Pt(10)
run.font.name = 'Consolas'
shading = OxmlElement('w:shd')
shading.set(qn('w:fill'), 'F5F5F5')
shading.set(qn('w:val'), 'clear')
p.paragraph_format.element.get_or_add_pPr().append(shading)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# 附录：快速参考
# ══════════════════════════════════════════════════════════════
doc.add_heading('附录：快速参考', level=1)

doc.add_heading('如果遇到问题', level=2)
issues = [
    ('构建失败', '发送：请检查 npm run build 的错误日志并修复'),
    ('部署失败', '发送：请先运行 setup_cloudflare_api_key 配置 API Token，然后重新部署'),
    ('页面空白', '发送：请检查浏览器控制台错误，用 PlaywrightConsoleCapture 抓取日志'),
    ('数据库错误', '发送：请检查 D1 查询日志，确认 migration 是否已执行'),
    ('GitHub Push 失败', '发送：请先运行 setup_github_environment 配置认证'),
]
for title, solution in issues:
    p = doc.add_paragraph()
    run = p.add_run(f'{title}: ')
    run.font.bold = True
    p.add_run(solution)

doc.add_heading('关键命令速查', level=2)
commands = [
    ('本地构建', 'cd /home/user/webapp && npm run build'),
    ('本地启动', 'cd /home/user/webapp && pm2 start ecosystem.config.cjs'),
    ('查看日志', 'pm2 logs zhongliu-connect --nostream'),
    ('部署生产', 'cd /home/user/webapp && npx wrangler pages deploy dist --project-name zhongliu-connect --branch main'),
    ('数据库迁移(本地)', 'cd /home/user/webapp && npx wrangler d1 migrations apply zhongliu-production --local'),
    ('数据库迁移(生产)', 'cd /home/user/webapp && npx wrangler d1 migrations apply zhongliu-production'),
    ('性能测试', 'curl -s -o /dev/null -w "TTFB: %{time_starttransfer}s" https://zhongliutong.net/'),
]
for label, cmd in commands:
    p = doc.add_paragraph()
    run = p.add_run(f'{label}: ')
    run.font.bold = True
    run2 = p.add_run(cmd)
    run2.font.name = 'Consolas'
    run2.font.size = Pt(9)

doc.add_heading('预期最终性能目标', level=2)
targets = [
    '所有页面 TTFB < 200ms',
    '页面 HTML < 100KB',
    'SPA 路由切换 < 300ms',
    'API 响应 < 100ms',
    '首屏可交互 (TTI) < 1.5s',
]
for t in targets:
    doc.add_paragraph(t, style='List Bullet')

# ── 保存 ──
output_path = '/home/user/webapp/zhongliu_performance_optimization_prompts.docx'
doc.save(output_path)
print(f'文档已生成: {output_path}')
